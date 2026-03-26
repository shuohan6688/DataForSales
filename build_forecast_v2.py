import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── 色定義 ──
C_TITLE      = "1F4E79"
C_HDR_MONTH  = "2E75B6"
C_HDR_GRAY   = "404040"
C_INPUT      = "FFF2CC"   # 黄：手入力
C_ACTUAL     = "DDEBF7"   # 青：実績（参照）
C_CALC       = "E2EFDA"   # 緑：自動計算
C_INV        = "BDD7EE"   # 濃青：在庫
C_DOS        = "FCE4D6"   # オレンジ：DOS
C_TOTAL      = "D9D9D9"   # グレー：合計
C_SECTION    = "595959"

MONTHS = ["1月","2月","3月","4月","5月","6月","7月","8月","9月","10月","11月","12月"]
DAYS   = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]

def hdr(cell, text, bg=C_TITLE, fc="FFFFFF", bold=True, sz=9, wrap=False):
    cell.value = text
    cell.font = Font(bold=bold, color=fc, size=sz)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def inp(cell, val=0, fmt='#,##0'):
    cell.value = val
    cell.fill = PatternFill("solid", fgColor=C_INPUT)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.number_format = fmt

def calc(cell, formula, fmt='#,##0', bold=False, bg=C_CALC):
    cell.value = formula
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.number_format = fmt
    if bold:
        cell.font = Font(bold=True)

def thin():
    s = Side(border_style="thin", color="BBBBBB")
    return Border(left=s, right=s, top=s, bottom=s)

def box(ws, r1, c1, r2, c2):
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            ws.cell(r, c).border = thin()

def section_label(ws, row, col, text, span=1):
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col+span-1)
    c = ws.cell(row, col, text)
    c.font = Font(bold=True, color="FFFFFF", size=9)
    c.fill = PatternFill("solid", fgColor=C_SECTION)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18

INP_SHEET = "①入力データ"

# ════════════════════════════════════════════════
# SHEET 1 : 入力データ
# ════════════════════════════════════════════════
ws1 = wb.active
ws1.title = INP_SHEET
ws1.sheet_view.showGridLines = False

# 列幅
col_w = {"A":16, "B":14, "C":14, "D":14, "E":13}
for col, w in col_w.items():
    ws1.column_dimensions[col].width = w

# タイトル
ws1.merge_cells("A1:E1")
c = ws1["A1"]
c.value = "在庫金額・DOS 予測モデル  ―  入力データ"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=C_TITLE)
c.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 28

# 凡例
ws1.merge_cells("A2:E2")
c = ws1["A2"]
c.value = "🟡 黄色 = 手入力   🔵 青色 = 実績（参照のみ）   🟢 緑色 = 自動計算"
c.font = Font(size=8, italic=True)
ws1.row_dimensions[2].height = 14

# ヘッダー行（共通）
def section_header(ws, row, label):
    section_label(ws, row, 1, label, span=5)

def col_headers(ws, row):
    ws.row_dimensions[row].height = 18
    for i, h in enumerate(["月", "CH1 店舗", "CH2 ROBO", "CH3 B2B", "合計"], 1):
        bg = C_HDR_GRAY if i == 1 else C_HDR_MONTH
        hdr(ws.cell(row, i, h), h, bg=bg)

# ─── Section 1: Sell Out 販売金額 ───
R = 4
section_header(ws1, R, "【Sell Out 販売金額】　単位：円"); R += 1
col_headers(ws1, R); R += 1
SALES_START = R
for mi in range(12):
    r = SALES_START + mi
    ws1.row_dimensions[r].height = 17
    tag = "実績" if mi < 2 else "目標"
    ws1.cell(r, 1, f"{MONTHS[mi]}（{tag}）").alignment = Alignment(horizontal="center", vertical="center")
    for col in [2, 3, 4]:
        c = ws1.cell(r, col)
        inp(c)
        if mi < 2:
            c.fill = PatternFill("solid", fgColor=C_ACTUAL)
    calc(ws1.cell(r, 5), f"=B{r}+C{r}+D{r}", bold=True, bg=C_TOTAL)
box(ws1, SALES_START, 1, SALES_START+11, 5)

# ─── Section 2: 仕入入庫金額 ───
R = SALES_START + 13
section_header(ws1, R, "【仕入入庫金額（채购到货）】　単位：円"); R += 1
col_headers(ws1, R); R += 1
PURCH_START = R
for mi in range(12):
    r = PURCH_START + mi
    ws1.row_dimensions[r].height = 17
    tag = "実績" if mi < 2 else "計画"
    ws1.cell(r, 1, f"{MONTHS[mi]}（{tag}）").alignment = Alignment(horizontal="center", vertical="center")
    for col in [2, 3, 4]:
        c = ws1.cell(r, col)
        inp(c)
        if mi < 2:
            c.fill = PatternFill("solid", fgColor=C_ACTUAL)
    calc(ws1.cell(r, 5), f"=B{r}+C{r}+D{r}", bold=True, bg=C_TOTAL)
box(ws1, PURCH_START, 1, PURCH_START+11, 5)

# ─── Section 3: 月末総在庫（実績）───
R = PURCH_START + 13
section_header(ws1, R, "【月末総在庫 実績】　1月末・2月末のみ入力　単位：円"); R += 1
col_headers(ws1, R); R += 1
INV_START = R
for mi in range(2):
    r = INV_START + mi
    ws1.row_dimensions[r].height = 17
    ws1.cell(r, 1, f"{MONTHS[mi]}末（実績）").alignment = Alignment(horizontal="center", vertical="center")
    for col in [2, 3, 4]:
        c = ws1.cell(r, col)
        inp(c)
        c.fill = PatternFill("solid", fgColor=C_ACTUAL)
    calc(ws1.cell(r, 5), f"=B{r}+C{r}+D{r}", bold=True, bg=C_TOTAL)
box(ws1, INV_START, 1, INV_START+1, 5)

# ─── Section 4: 参考 店舗数・ROBO台数 ───
R = INV_START + 4
section_header(ws1, R, "【参考】店舗数・ROBO台数（DOS分析の補助用）"); R += 1
ws1.row_dimensions[R].height = 18
hdr(ws1.cell(R, 1), "月", bg=C_HDR_GRAY)
hdr(ws1.cell(R, 2), "CH1 店舗数（月末）", bg=C_HDR_MONTH)
hdr(ws1.cell(R, 3), "CH2 ROBO台数（月末）", bg=C_HDR_MONTH)
hdr(ws1.cell(R, 4), "備考", bg=C_HDR_GRAY)
ws1.column_dimensions["D"].width = 20
R += 1
STORE_START = R
for mi in range(12):
    r = STORE_START + mi
    ws1.row_dimensions[r].height = 17
    ws1.cell(r, 1, MONTHS[mi]).alignment = Alignment(horizontal="center", vertical="center")
    inp(ws1.cell(r, 2), 0, '0')
    inp(ws1.cell(r, 3), 0, '0')
    ws1.cell(r, 4, "").alignment = Alignment(horizontal="left")
box(ws1, STORE_START, 1, STORE_START+11, 4)

# メタデータ保存
META = {
    "sales_start": SALES_START,
    "purch_start": PURCH_START,
    "inv_start":   INV_START,
    "store_start": STORE_START,
}

# ════════════════════════════════════════════════
# SHEET 2 : 予測計算（チャネル別）
# ════════════════════════════════════════════════
ws2 = wb.create_sheet("②予測計算")
ws2.sheet_view.showGridLines = False

# 列構成：A=指標, B=1月 ... M=12月
ws2.column_dimensions["A"].width = 22
for ci in range(2, 15):
    ws2.column_dimensions[get_column_letter(ci)].width = 13

# タイトル
ws2.merge_cells("A1:M1")
c = ws2["A1"]
c.value = "在庫金額・DOS 月次予測計算シート"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=C_TITLE)
c.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 28

ws2.merge_cells("A2:M2")
c = ws2["A2"]
c.value = "数式は自動設定済み。1月・2月は実績を参照。3月以降は入力データシートの計画値から自動計算。"
c.font = Font(size=8, italic=True)
ws2.row_dimensions[2].height = 14

# 月ヘッダー行
ws2.row_dimensions[3].height = 20
ws2.cell(3, 1, "指標").alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(3, 1).fill = PatternFill("solid", fgColor=C_HDR_GRAY)
ws2.cell(3, 1).font = Font(bold=True, color="FFFFFF", size=9)
for mi, month in enumerate(MONTHS):
    c = ws2.cell(3, 2+mi)
    c.value = month
    c.font = Font(bold=True, color="FFFFFF", size=9)
    c.fill = PatternFill("solid", fgColor=C_HDR_MONTH)
    c.alignment = Alignment(horizontal="center", vertical="center")

S  = INP_SHEET
SS = META["sales_start"]
PS = META["purch_start"]
IS = META["inv_start"]

def month_col(mi):  # mi=0→B, 1→C ...
    return get_column_letter(2 + mi)

def build_channel_block(ws, start_row, ch_name, ch_col_in_input, bg_section):
    """
    ch_col_in_input: 2=CH1, 3=CH2, 4=CH3 (column B/C/D in input sheet)
    Returns: dict of row numbers for key metrics
    """
    R = start_row

    # セクションタイトル
    ws.merge_cells(start_row=R, start_column=1, end_row=R, end_column=13)
    c = ws.cell(R, 1, f"  {ch_name}")
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=bg_section)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[R].height = 20
    R += 1

    # 指標ラベルと行番号の管理
    row_labels = [
        ("Sell Out 販売金額",  "sell"),
        ("仕入入庫金額",       "purch"),
        ("月末在庫金額",       "inv"),
        ("翌月 日別Sell Out",  "daily"),
        ("DOS（日）",          "dos"),
    ]
    rows = {}
    for label, key in row_labels:
        rows[key] = R
        ws.row_dimensions[R].height = 17
        c = ws.cell(R, 1, label)
        c.font = Font(size=9, bold=(key in ["inv", "dos"]))
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.fill = PatternFill("solid", fgColor="F7F7F7")
        R += 1

    # 数式を埋める
    inv_row   = rows["inv"]
    sell_row  = rows["sell"]
    purch_row = rows["purch"]
    daily_row = rows["daily"]
    dos_row   = rows["dos"]

    for mi in range(12):
        col = month_col(mi)
        sr = SS + mi   # 販売行
        pr = PS + mi   # 仕入行
        ir = IS + mi   # 在庫実績行（IS, IS+1 = Jan, Feb）

        # ── Sell Out ──
        c = ws.cell(sell_row, 2+mi)
        c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{sr}"
        c.number_format = '#,##0'
        c.fill = PatternFill("solid", fgColor=C_ACTUAL if mi < 2 else C_CALC)
        c.alignment = Alignment(horizontal="right", vertical="center")

        # ── 仕入入庫 ──
        c = ws.cell(purch_row, 2+mi)
        c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{pr}"
        c.number_format = '#,##0'
        c.fill = PatternFill("solid", fgColor=C_ACTUAL if mi < 2 else C_INPUT)
        c.alignment = Alignment(horizontal="right", vertical="center")

        # ── 月末在庫 ──
        c = ws.cell(inv_row, 2+mi)
        c.font = Font(bold=True)
        c.number_format = '#,##0'
        c.fill = PatternFill("solid", fgColor=C_INV)
        c.alignment = Alignment(horizontal="right", vertical="center")
        if mi == 0:
            # 1月：実績参照
            c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{IS}"
        elif mi == 1:
            # 2月：実績参照
            c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{IS+1}"
        else:
            # 3月以降：前月末在庫 + 仕入 - 販売
            prev_col = month_col(mi-1)
            c.value = (f"={prev_col}{inv_row}"
                       f"+{col}{purch_row}"
                       f"-{col}{sell_row}")

        # ── 翌月 日別Sell Out ──
        c = ws.cell(daily_row, 2+mi)
        c.number_format = '#,##0'
        c.fill = PatternFill("solid", fgColor=C_CALC)
        c.alignment = Alignment(horizontal="right", vertical="center")
        if mi < 11:
            next_sr = SS + mi + 1
            next_days = DAYS[mi+1]
            c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{next_sr}/{next_days}"
        else:
            # 12月：12月自身の日販
            c.value = f"='{S}'!{get_column_letter(ch_col_in_input)}{SS+11}/{DAYS[11]}"

        # ── DOS ──
        c = ws.cell(dos_row, 2+mi)
        c.value = f"=IF({col}{daily_row}>0, {col}{inv_row}/{col}{daily_row}, 0)"
        c.number_format = '0.0'
        c.fill = PatternFill("solid", fgColor=C_DOS)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(bold=True)

    box(ws, rows["sell"], 1, rows["dos"], 13)
    return rows, R + 1

# 各チャネルブロック
CH_CONFIGS = [
    ("CH1 店舗", 2, "1F4E79"),
    ("CH2 ROBO", 3, "375623"),
    ("CH3 B2B",  4, "7B3F00"),
]

current_row = 5
ch_rows = {}
for ch_name, ch_col, bg in CH_CONFIGS:
    rows, current_row = build_channel_block(ws2, current_row, ch_name, ch_col, bg)
    ch_rows[ch_name] = rows
    current_row += 1  # 空行

# ── 合計ブロック ──
ws2.merge_cells(start_row=current_row, start_column=1,
                end_row=current_row, end_column=13)
c = ws2.cell(current_row, 1, "  合計（全チャネル）")
c.font = Font(bold=True, color="FFFFFF", size=10)
c.fill = PatternFill("solid", fgColor="1F3864")
c.alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[current_row].height = 20
current_row += 1

total_labels = [
    ("合計 Sell Out 販売", "sell"),
    ("合計 仕入入庫",      "purch"),
    ("合計 月末在庫",      "inv"),
    ("翌月 日別 合計Sell Out", "daily"),
    ("合計 DOS（日）",     "dos"),
]
total_rows = {}
for label, key in total_labels:
    total_rows[key] = current_row
    c = ws2.cell(current_row, 1, label)
    c.font = Font(size=9, bold=(key in ["inv", "dos"]))
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.fill = PatternFill("solid", fgColor="F0F0F0")
    ws2.row_dimensions[current_row].height = 17
    current_row += 1

ch1_r = ch_rows["CH1 店舗"]
ch2_r = ch_rows["CH2 ROBO"]
ch3_r = ch_rows["CH3 B2B"]

for mi in range(12):
    col = month_col(mi)

    # 合計 Sell Out
    c = ws2.cell(total_rows["sell"], 2+mi)
    c.value = (f"={col}{ch1_r['sell']}"
               f"+{col}{ch2_r['sell']}"
               f"+{col}{ch3_r['sell']}")
    c.number_format = '#,##0'; c.fill = PatternFill("solid", fgColor=C_ACTUAL if mi < 2 else C_CALC)
    c.alignment = Alignment(horizontal="right", vertical="center")

    # 合計 仕入
    c = ws2.cell(total_rows["purch"], 2+mi)
    c.value = (f"={col}{ch1_r['purch']}"
               f"+{col}{ch2_r['purch']}"
               f"+{col}{ch3_r['purch']}")
    c.number_format = '#,##0'; c.fill = PatternFill("solid", fgColor=C_ACTUAL if mi < 2 else C_INPUT)
    c.alignment = Alignment(horizontal="right", vertical="center")

    # 合計 月末在庫
    c = ws2.cell(total_rows["inv"], 2+mi)
    c.value = (f"={col}{ch1_r['inv']}"
               f"+{col}{ch2_r['inv']}"
               f"+{col}{ch3_r['inv']}")
    c.number_format = '#,##0'; c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="9DC3E6")
    c.alignment = Alignment(horizontal="right", vertical="center")

    # 合計 翌月日販
    if mi < 11:
        next_sr = SS + mi + 1
        next_days = DAYS[mi+1]
        daily_formula = f"='{S}'!E{next_sr}/{next_days}"
    else:
        daily_formula = f"='{S}'!E{SS+11}/{DAYS[11]}"
    c = ws2.cell(total_rows["daily"], 2+mi)
    c.value = daily_formula
    c.number_format = '#,##0'; c.fill = PatternFill("solid", fgColor=C_CALC)
    c.alignment = Alignment(horizontal="right", vertical="center")

    # 合計 DOS
    c = ws2.cell(total_rows["dos"], 2+mi)
    c.value = (f"=IF({col}{total_rows['daily']}>0,"
               f"{col}{total_rows['inv']}/{col}{total_rows['daily']},0)")
    c.number_format = '0.0'; c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="C55A11")
    c.alignment = Alignment(horizontal="center", vertical="center")

box(ws2, total_rows["sell"], 1, total_rows["dos"], 13)

# ════════════════════════════════════════════════
# SHEET 3 : サマリダッシュボード
# ════════════════════════════════════════════════
ws3 = wb.create_sheet("③サマリ")
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 20
for ci in range(2, 15):
    ws3.column_dimensions[get_column_letter(ci)].width = 11

ws3.merge_cells("A1:M1")
c = ws3["A1"]
c.value = "サマリダッシュボード　―　月次在庫金額・DOS 推移"
c.font = Font(bold=True, size=13, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=C_TITLE)
c.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 28

# 月ヘッダー
ws3.row_dimensions[3].height = 20
hdr(ws3.cell(3, 1), "指標", bg=C_HDR_GRAY)
for mi in range(12):
    hdr(ws3.cell(3, 2+mi), MONTHS[mi], bg=C_HDR_MONTH)

CALC = "②予測計算"
summary_rows = [
    ("─── Sell Out 販売 ───",  None, None, "44546A"),
    ("CH1 店舗",               ch1_r["sell"],   C_CALC,   None),
    ("CH2 ROBO",               ch2_r["sell"],   C_CALC,   None),
    ("CH3 B2B",                ch3_r["sell"],   C_CALC,   None),
    ("合計 Sell Out",          total_rows["sell"], C_TOTAL, None),
    ("─── 月末在庫金額 ───",   None, None, "44546A"),
    ("CH1 店舗",               ch1_r["inv"],    C_INV,    None),
    ("CH2 ROBO",               ch2_r["inv"],    C_INV,    None),
    ("CH3 B2B",                ch3_r["inv"],    C_INV,    None),
    ("合計 在庫",              total_rows["inv"], "9DC3E6", None),
    ("─── DOS（日）───",       None, None, "44546A"),
    ("CH1 DOS",                ch1_r["dos"],    C_DOS,    None),
    ("CH2 DOS",                ch2_r["dos"],    C_DOS,    None),
    ("CH3 DOS",                ch3_r["dos"],    C_DOS,    None),
    ("合計 DOS",               total_rows["dos"], "C55A11", None),
]

sr = 4
for label, src_row, bg, section_bg in summary_rows:
    ws3.row_dimensions[sr].height = 17
    if src_row is None:
        ws3.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=13)
        c = ws3.cell(sr, 1, f"  {label}")
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=section_bg)
        c.alignment = Alignment(horizontal="left", vertical="center")
    else:
        c = ws3.cell(sr, 1, label)
        c.font = Font(size=9)
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.fill = PatternFill("solid", fgColor="F7F7F7")
        is_dos = (bg in [C_DOS, "C55A11"])
        is_dos_total = (bg == "C55A11")
        for mi in range(12):
            col = month_col(mi)
            ref_cell = f"'{CALC}'!{col}{src_row}"
            c2 = ws3.cell(sr, 2+mi, f"={ref_cell}")
            c2.number_format = '0.0' if is_dos else '#,##0'
            c2.fill = PatternFill("solid", fgColor=bg)
            c2.alignment = Alignment(
                horizontal="center" if is_dos else "right",
                vertical="center")
            if is_dos_total or label.startswith("合計"):
                c2.font = Font(bold=True,
                               color="FFFFFF" if is_dos_total else "000000")
    sr += 1

box(ws3, 3, 1, sr-1, 13)

# 保存
out_path = "/home/user/DataForSales/在庫DOS予測モデル_v2.xlsx"
wb.save(out_path)
print(f"保存完了: {out_path}")
