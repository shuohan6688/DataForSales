"""
CrossUse（他社IP×自社IP）セグメント分析
=========================================
注文ID（カート）単位で、以下の3セグメントに分類して集計します：
  - 他社IPのみ  : 他社IPコラボ商品だけ購入
  - クロス購買  : 同一カートで他社IP + 自社IP 両方購入
  - 自社IPのみ  : 自社IP商品だけ購入

※ shoppingbag は集計から除外します

【必要なファイル】
  1枚のExcel/CSVシートに以下の列が含まれること：
    order_id, customer_id, Product_id, quantity, sales, order_date, Separate

【使い方】
  1. INPUT_FILE に実データのパスを設定
  2. python analyze_crossuse.py
  3. results/crossuse_analysis.xlsx に結果が出力される
"""

import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── ファイルパス設定 ───────────────────────────────────
INPUT_FILE  = "sample_data/orders_sample.xlsx"   # ← 実データのパスに変更
OUTPUT_FILE = "results/crossuse_analysis.xlsx"

# ── 列名設定（実データの列名に合わせて変更） ──────────────
ORDER_ID_COL = "order_id"
TYPE_COL     = "Separate"    # 商品区分列
SALES_COL    = "sales"       # 単価列
QTY_COL      = "quantity"    # 数量列

# ── Separate列の値 ───────────────────────────────────
COLLAB_VALUE   = "他社IP"
OWN_IP_VALUE   = "自社IP"
EXCLUDE_VALUES = ["shoppingbag"]   # 集計除外する値


# ────────────────────────────────────────────────
# 1. データ読み込み
# ────────────────────────────────────────────────
def load_data(input_path: str) -> pd.DataFrame:
    p = Path(input_path)
    df = pd.read_excel(p) if p.suffix in (".xlsx", ".xls") else pd.read_csv(p)
    print(f"[読込] {len(df):,} 行")

    # Separate列の値の確認
    print(f"[確認] Separate列の値: {df[TYPE_COL].value_counts().to_dict()}")
    return df


# ────────────────────────────────────────────────
# 2. 前処理（shoppingbag除外）
# ────────────────────────────────────────────────
def preprocess(df: pd.DataFrame) -> pd.DataFrame:
    before = len(df)
    raw_sales = (df[SALES_COL] * df[QTY_COL]).sum()
    print(f"[確認] 除外前の売上合計（単価×数量）: {raw_sales:,.0f}")

    df = df[~df[TYPE_COL].isin(EXCLUDE_VALUES)].copy()
    excluded = before - len(df)

    # 売上金額列を計算して追加
    df["line_total"] = df[SALES_COL] * df[QTY_COL]
    after_sales = df["line_total"].sum()
    print(f"[除外] shoppingbag: {excluded:,} 行 → 残り {len(df):,} 行")
    print(f"[確認] 除外後の売上合計: {after_sales:,.0f}")
    return df


# ────────────────────────────────────────────────
# 3. 注文ID単位でセグメント分類
# ────────────────────────────────────────────────
def classify_orders(df: pd.DataFrame) -> pd.DataFrame:
    grouped = df.groupby(ORDER_ID_COL)[TYPE_COL]
    order_flags = pd.DataFrame({
        "has_collab": grouped.apply(lambda t: (t == COLLAB_VALUE).any()),
        "has_own_ip": grouped.apply(lambda t: (t == OWN_IP_VALUE).any()),
    }).reset_index()

    # 注文ID単位の売上合計（単価×数量）を結合
    order_sales = df.groupby(ORDER_ID_COL)["line_total"].sum().reset_index()
    order_flags = order_flags.merge(order_sales, on=ORDER_ID_COL, how="left")
    order_flags.rename(columns={"line_total": SALES_COL}, inplace=True)

    def segment(row):
        if row["has_collab"] and row["has_own_ip"]:
            return "cross"
        elif row["has_collab"]:
            return "collab_only"
        else:
            return "own_ip_only"

    order_flags["segment"] = order_flags.apply(segment, axis=1)
    return order_flags


# ────────────────────────────────────────────────
# 4. 集計
# ────────────────────────────────────────────────
def summarize(order_flags: pd.DataFrame) -> pd.DataFrame:
    total_orders = len(order_flags)
    total_sales  = order_flags[SALES_COL].sum()

    agg = order_flags.groupby("segment").agg(
        注文数=(ORDER_ID_COL, "count"),
        売上合計=(SALES_COL, "sum"),
    )

    rows = []
    for key, label in [
        ("collab_only", "他社IP（コラボ）のみ"),
        ("cross",       "他社IP + 自社IP（クロス購買）"),
        ("own_ip_only", "自社IPのみ"),
    ]:
        n = int(agg.loc[key, "注文数"]) if key in agg.index else 0
        s = float(agg.loc[key, "売上合計"]) if key in agg.index else 0.0
        rows.append({
            "segment":       key,
            "セグメント名":  label,
            "注文数":        n,
            "注文構成比(%)": round(n / total_orders * 100, 1),
            "売上合計":      s,
            "売上構成比(%)": round(s / total_sales * 100, 1) if total_sales else 0.0,
        })

    summary = pd.DataFrame(rows)

    total_row = pd.DataFrame([{
        "segment":       "total",
        "セグメント名":  "合計",
        "注文数":        total_orders,
        "注文構成比(%)": 100.0,
        "売上合計":      total_sales,
        "売上構成比(%)": 100.0,
    }])
    summary = pd.concat([summary, total_row], ignore_index=True)

    print("\n=== CrossUseセグメント集計結果 ===")
    print(summary[["セグメント名", "注文数", "注文構成比(%)", "売上合計", "売上構成比(%)"]].to_string(index=False))
    print()
    return summary


# ────────────────────────────────────────────────
# 5. Excel出力
# ────────────────────────────────────────────────
BLUE_HEADER = "1F4E79"
COLORS = {
    "collab_only": "F4B942",   # オレンジ
    "cross":       "70AD47",   # グリーン
    "own_ip_only": "4472C4",   # ブルー
    "total":       "595959",   # グレー
}


def _thin_border():
    s = Side(border_style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)


def export_excel(summary: pd.DataFrame, order_flags: pd.DataFrame, output_path: str):
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # ── シート1: サマリ ──────────────────────────────────
    ws = wb.active
    ws.title = "セグメント集計"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "CrossUse分析 ― 他社IP × 自社IP セグメント集計（カート単位）"
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 注釈
    ws.merge_cells("A2:C2")
    c = ws["A2"]
    c.value = "※ shoppingbag は集計除外済み　／　1カート（order_id）を1取引として集計"
    c.font = Font(size=9, italic=True, color="666666")
    c.alignment = Alignment(horizontal="left")
    ws.row_dimensions[2].height = 16

    ws.merge_cells("A1:E1")
    ws["A1"].value = "CrossUse分析 ― 他社IP × 自社IP セグメント集計（カート単位）"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE_HEADER)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:E2")

    headers    = ["セグメント名", "注文数（カート）", "注文構成比", "売上合計", "売上構成比"]
    col_widths = [36, 18, 14, 18, 14]
    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(3, i, h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20

    for row_idx, row in summary.iterrows():
        r = row_idx + 4
        seg_key  = row["segment"]
        seg_name = row["セグメント名"]
        color    = COLORS.get(seg_key, "CCCCCC")
        is_total = (seg_key == "total")
        bg       = "DDDDDD" if is_total else "F5F5F5"

        c = ws.cell(r, 1, seg_name)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        c = ws.cell(r, 2, row["注文数"])
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")
        c.fill = PatternFill("solid", fgColor=bg)
        if is_total: c.font = Font(bold=True)

        c = ws.cell(r, 3, row["注文構成比(%)"] / 100)
        c.number_format = '0.0%'
        c.alignment = Alignment(horizontal="right")
        c.fill = PatternFill("solid", fgColor=bg)
        if is_total: c.font = Font(bold=True)

        c = ws.cell(r, 4, row["売上合計"])
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")
        c.fill = PatternFill("solid", fgColor=bg)
        if is_total: c.font = Font(bold=True)

        c = ws.cell(r, 5, row["売上構成比(%)"] / 100)
        c.number_format = '0.0%'
        c.alignment = Alignment(horizontal="right")
        c.fill = PatternFill("solid", fgColor=bg)
        if is_total: c.font = Font(bold=True)

        ws.row_dimensions[r].height = 22
        for col in range(1, 6):
            ws.cell(r, col).border = _thin_border()

    for col in range(1, 6):
        ws.cell(3, col).border = _thin_border()

    # ── シート2: 注文IDリスト ────────────────────────────
    ws2 = wb.create_sheet("注文セグメント一覧")
    ws2.sheet_view.showGridLines = False

    ws2.merge_cells("A1:D1")
    c = ws2["A1"]
    c.value = "注文ID別 セグメント分類一覧（shoppingbag除外済み）"
    c.font = Font(bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 24

    seg_label_map = {
        "collab_only": "他社IP（コラボ）のみ",
        "cross":       "他社IP + 自社IP（クロス購買）",
        "own_ip_only": "自社IPのみ",
    }
    detail_headers = [ORDER_ID_COL, "他社IP含む", "自社IP含む", "売上合計", "セグメント"]
    detail_widths  = [20, 14, 14, 16, 34]
    for i, (h, w) in enumerate(zip(detail_headers, detail_widths), 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
        c = ws2.cell(2, i, h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    for row_idx, row in order_flags.reset_index(drop=True).iterrows():
        r = row_idx + 3
        seg_key   = row["segment"]
        seg_label = seg_label_map[seg_key]
        color     = COLORS[seg_key]

        ws2.cell(r, 1, row[ORDER_ID_COL]).alignment = Alignment(horizontal="center")
        ws2.cell(r, 2, "○" if row["has_collab"] else "―").alignment = Alignment(horizontal="center")
        ws2.cell(r, 3, "○" if row["has_own_ip"] else "―").alignment = Alignment(horizontal="center")

        c = ws2.cell(r, 4, row[SALES_COL])
        c.number_format = '#,##0'
        c.alignment = Alignment(horizontal="right")

        c = ws2.cell(r, 5, seg_label)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="left", indent=1)
        ws2.row_dimensions[r].height = 16

        for col in range(1, 6):
            ws2.cell(r, col).border = _thin_border()

    wb.save(output_path)
    print(f"[出力] {output_path}")


# ────────────────────────────────────────────────
# メイン実行
# ────────────────────────────────────────────────
if __name__ == "__main__":
    df          = load_data(INPUT_FILE)
    df          = preprocess(df)
    order_flags = classify_orders(df)
    summary     = summarize(order_flags)
    export_excel(summary, order_flags, OUTPUT_FILE)
